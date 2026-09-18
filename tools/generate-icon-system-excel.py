#!/usr/bin/env python3
"""
生成 06_ICON_DESIGN_AND_RANKING_SYSTEM.xlsx
记录图标量化评级模型 (ISO 7001)、打分权重、候选对比与选定成果
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def generate_icon_excel():
    wb = openpyxl.Workbook()
    
    # 工作表 1：图标筛选标准与权重模型
    ws1 = wb.active
    ws1.title = "01_评估标准与权重模型"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    border_style = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)

    headers1 = ["维度编号", "评估维度 (Evaluation Dimensions)", "权重 (Weight)", "参考标准 (International Standard)", "量化考量细则 (Criteria Details)"]
    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c_data = [
        ["CR-01", "跨文化/全球通用认知度 (Global Recognizability)", "35%", "ISO 7001 / IEC 60417", "符号必须具备超越语言的物理直觉，杜绝特定区域本土土俗隐喻，欧亚美用户一眼即懂"],
        ["CR-02", "机械与智能硬件工程契合度 (Hardware Engineering Fit)", "25%", "DIN / ANSI Hardware Blueprint", "杜绝互联网安全'卡通挂锁'，必须体现精密锁芯截面、卡扣传动销或工程装配物理特征"],
        ["CR-03", "小尺寸高辨识与几何平衡度 (Visual Scalability)", "20%", "SVG Grid 24×24px Baseline", "在 15px ~ 24px 极小视网膜屏幕下依然线条清晰，不糊点，视觉重心绝对居中"],
        ["CR-04", "高级工业质感与去俗套化 (Aesthetic Elegance)", "20%", "Industrial Minimalist", "采用深空灰/钛金蓝微渐变与精密线条，告别俗气亮蓝与厚重块面，呈现硬核知识库气质"]
    ]
    for r in c_data:
        ws1.append(r)

    # 工作表 2：全站图标候选比对与综合打分表
    ws2 = wb.create_sheet(title="02_全站图标打分与入选清单")
    headers2 = ["应用位置", "导航/品牌项", "候选方案类型", "全球认知度 (35%)", "工程契合度 (25%)", "小尺寸清晰度 (20%)", "工业美学度 (20%)", "综合加权分 (100)", "评审决策", "设计理念与物理特征"]
    ws2.append(headers2)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    items_data = [
        ["左上角 Brand", "全局品牌主标", "卡通纯色挂锁 (原版)", 88, 45, 70, 40, 64.05, "淘汰 (REJECTED)", "太像通用网页 SSL 安全锁，无机械锁具工业感，俗气卡通"],
        ["左上角 Brand", "全局品牌主标", "极简单线钥匙", 85, 60, 80, 75, 75.75, "备选 (CANDIDATE)", "偏向软件权限或房产中介，缺乏门锁与改装工程厚重度"],
        ["左上角 Brand", "全局品牌主标", "欧标水滴双向锁芯+高精转心", 96, 98, 95, 96, 96.30, "入选 (SELECTED)", "以全球最主流 DIN 锁芯外形为底，内嵌钛蓝渐变与 1:1 钥匙孔传动轴，高级硬核"],
        ["Menu 菜单", "锁型总览 (Catalog)", "通用文件夹/列表", 82, 65, 85, 70, 75.95, "淘汰 (REJECTED)", "太普通，偏办公文档感"],
        ["Menu 菜单", "锁型总览 (Catalog)", "精密锁舌装配框 (Lock Assembly)", 95, 96, 94, 95, 95.05, "入选 (SELECTED)", "象征 5 大工业板块与锁体矩形面阵，清晰传达机械门锁品类"],
        ["Menu 菜单", "工程实录 (Field Cases)", "通俗单反照相机", 85, 70, 85, 75, 79.25, "淘汰 (REJECTED)", "容易被误认为摄影作品图库，非工程现场实拍"],
        ["Menu 菜单", "工程实录 (Field Cases)", "游标卡尺与质检十字 (Calipers / QA)", 94, 98, 92, 96, 95.00, "入选 (SELECTED)", "深度契合施工现场装配、间隙测量与一线实测质检工程感"],
        ["Menu 菜单", "转接工具 (Adapters & BOM)", "单只普通呆扳手", 85, 80, 85, 75, 81.75, "淘汰 (REJECTED)", "通用修车维修工具，未突出轴套变径与转接机构"],
        ["Menu 菜单", "转接工具 (Adapters & BOM)", "精密传动轴套与齿轮 (Coupler & Gear)", 93, 97, 95, 94, 94.60, "入选 (SELECTED)", "完美呼应 7mm转8mm 轴套、偏心转轴及五金加装 BOM 配件"],
        ["Menu 菜单", "避坑实录 (Field Pitfalls)", "普通感叹号警示三角", 90, 75, 88, 80, 83.85, "淘汰 (REJECTED)", "通用网页报错图标，缺乏机械碰撞受阻具象感"],
        ["Menu 菜单", "避坑实录 (Field Pitfalls)", "错位剪切卡阻三角盾 (Mechanical Shear)", 96, 95, 94, 93, 94.75, "入选 (SELECTED)", "传递扣板错位、密封条挤压、电机过载等高危避坑警戒"],
        ["Menu 菜单", "工业索引 (Master Index)", "普通纯文本横线", 80, 60, 82, 70, 73.40, "淘汰 (REJECTED)", "缺乏矩阵与全球架构的专业体系感"],
        ["Menu 菜单", "工业索引 (Master Index)", "多维坐标图谱与蓝图立方 (Global Matrix)", 95, 96, 95, 95, 95.25, "入选 (SELECTED)", "代表 54 款锁型与 17 项国际标准的立体多维交叉索引图谱"]
    ]
    for r in items_data:
        ws2.append(r)

    # 自动列宽
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    out_path = os.path.join(ROOT, "docs", "06_ICON_DESIGN_AND_RANKING_SYSTEM.xlsx")
    wb.save(out_path)
    # 同步备份至用户主目录
    import shutil
    shutil.copy(out_path, "/home/user/06_ICON_DESIGN_AND_RANKING_SYSTEM.xlsx")
    print("Generated docs/06_ICON_DESIGN_AND_RANKING_SYSTEM.xlsx successfully!")

if __name__ == "__main__":
    generate_icon_excel()
