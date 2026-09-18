#!/usr/bin/env python3
"""
生成 08_HARDWARE_ADAPTERS_AND_BOM.xlsx
记录全量 8 款经过实物图纸对照与扭矩应力核实的转接五金标准清单
"""
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOM_JSON = os.path.join(ROOT, "content", "catalog", "adapters-bom.json")

def generate_adapters_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01_转接件全量工程核实BOM"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    border_style = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)

    headers = [
        "零件编号", "配件名称", "适用区域与标准", "解决的一线工程痛点", 
        "关键加工公差与尺寸", "推荐工程材质", "3D打印支持度", "核实状态与测试依据", "研发选型避坑警告"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    with open(BOM_JSON, "r", encoding="utf-8") as f:
        adapters = json.load(f)

    for a in adapters:
        ws.append([
            a.get("id"),
            a.get("name"),
            a.get("targetRegion"),
            a.get("problemSolved"),
            a.get("criticalTolerance"),
            a.get("materialRecommendation"),
            "支持 3D 打印快速打样" if a.get("diy3dPrintReady") else "严禁塑料件·必须机加工金属",
            a.get("verificationStatus"),
            a.get("notes")
        ])

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    out_path = os.path.join(ROOT, "docs", "08_HARDWARE_ADAPTERS_AND_BOM.xlsx")
    wb.save(out_path)
    import shutil
    shutil.copy(out_path, "/home/user/08_HARDWARE_ADAPTERS_AND_BOM.xlsx")
    print("Generated docs/08_HARDWARE_ADAPTERS_AND_BOM.xlsx successfully!")

if __name__ == "__main__":
    generate_adapters_excel()
