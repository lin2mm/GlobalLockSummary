#!/usr/bin/env python3
"""
重新生成 docs/03_GLOBAL_LOCK_DATA_INDEX.xlsx 及 docs/GLOBAL_LOCK_DATA_INDEX.xlsx
包含全站 54 款锁型的全维度工程数据、公差矩阵、二次修槽扩孔规范、电机功耗、防钻保险等级与门磁屏蔽指引
"""
import os
import json
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GALLERY_JSON = os.path.join(ROOT, "content", "catalog", "gallery.json")

def build_master_excel():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "01_全局工业索引与标准"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    border_style = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)

    ws1_headers = [
        "序号", "工业板块", "标准代码", "样本总数", "主流背距 (Backset)", 
        "中心距 (Centres)", "方轴孔径", "门厚区间", "电机工作峰值电流", "防钻防撬保险等级", "门磁感应间隙规范"
    ]
    ws1.append(ws1_headers)
    for col_idx in range(1, len(ws1_headers) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    divisions_data = [
        ["DIV-01", "北美工业板块 (North America)", "ANSI / BHMA A156", 10, "60/70mm 可调", "140mm (5-1/2\")", "1.6×4.8mm 扁平尾轴", "35-51mm", "堵转峰值 1.6A~2.2A (要求内阻≤120mΩ锂铁电池)", "ANSI/BHMA Grade 1 最高防盗级别", "间隙 12-18mm；金属门套需加 3mm 隔磁垫片"],
        ["DIV-02", "欧陆工业板块 (Continental Europe)", "DIN 18251 / EN 12209", 12, "55/65mm (窄框35-45mm)", "室内72mm / 入户92mm", "8×8mm (逃生9×9mm)", "38-65mm", "多点锁峰值 2.4A~2.8A (需 3C 高倍率放电锂包)", "SKG★★★ / VdS 2162 Class B (Allianz/AXA 认可)", "间隙 10-15mm；装甲门采用磁簧+陀螺仪双模检测"],
        ["DIV-03", "澳新及英国板块 (Oceania & UK)", "AS 4145 / BS 3621", 12, "40/60mm", "分体夜闩", "十字/偏心尾轴", "32-45mm", "压紧峰值 1.5A~2.0A (瞬态跌落电压≤0.3V)", "BS 3621 / Sold Secure Diamond 钻石级防盗", "间隙 8-14mm；外铁门建议偏置安装并预留5mm沉降"],
        ["DIV-04", "东南亚与东亚板块 (East & Southeast Asia)", "JIS A 1510 / SS 332", 12, "51/64mm", "独立/联动", "8×8mm 高精", "33-42mm", "极小阻尼峰值 1.2A~1.5A (常规电池续航超15个月)", "JIS A 1510 / 日本 CP 防盗认证 (防撬防钻≥10分)", "间隙 5-10mm；提供抗金属磁屏蔽微型磁铁组件"],
        ["DIV-05", "拉美工业板块 (Latin America)", "ABNT NBR 14913", 8, "40/45mm 极窄", "53/70mm", "8×8mm 宽公差", "30-35mm", "高摩擦峰值 2.2A~2.6A (需 500ms 快速过流熔断保护)", "ABNT NBR 14913 标准防撬防锯等级", "间隙 10-18mm；配备双孔螺丝紧固型 N52 强磁体"]
    ]
    for row in divisions_data:
        ws1.append(row)

    ws2 = wb.create_sheet(title="02_54款机械锁详细数据库")
    ws2_headers = [
        "锁型ID", "候选标识", "锁系ID", "中文名称", "英文名称", "工业板块",
        "市场保有率", "加装亲和度", "场景实景图路径", "机械结构图路径",
        "门缝公差要求", "标称电机扭矩", "垂直下沉耐受", "逃生安全规范", "租房友好评级",
        "钥匙胚槽型", "木槽二次扩孔修整指引", "对穿螺栓避线通道", "电机峰值电流与电源要求", "防盗安全与保险资质"
    ]
    ws2.append(ws2_headers)
    for col_idx in range(1, len(ws2_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    with open(GALLERY_JSON, "r", encoding="utf-8") as f:
        locks = json.load(f)

    for lk in locks:
        sel = lk.get("selectionScore", {})
        guide = lk.get("installationGuide", {})
        eng = lk.get("engineeringMatrix", {})
        egress = eng.get("egressCompliance", {})
        rental = eng.get("rentalOptimization", {})

        row_data = [
            lk.get("id"),
            lk.get("candidateId"),
            lk.get("familyId"),
            lk.get("title", {}).get("zh", ""),
            lk.get("title", {}).get("en", ""),
            lk.get("region"),
            sel.get("marketCoverage", "High") if isinstance(sel, dict) else "High",
            sel.get("retrofitAffinity", "Grade A"),
            lk.get("sceneImage", lk.get("image")),
            lk.get("productImage", lk.get("image")),
            guide.get("recommendedClearance", "≥ 3.0mm"),
            guide.get("requiredTorque", "≥ 1.5 N·m"),
            eng.get("saggingTolerance", "±2.0mm"),
            egress.get("standard", "Compliant"),
            rental.get("rating", "Grade A"),
            eng.get("keywaySpecification", "N/A"),
            eng.get("mortiseReworkGuide", "N/A"),
            eng.get("boltWireClearance", "N/A"),
            eng.get("electricalProfile", "N/A"),
            eng.get("securityCertification", "N/A")
        ]
        ws2.append(row_data)

    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    out_docs = os.path.join(ROOT, "docs", "03_GLOBAL_LOCK_DATA_INDEX.xlsx")
    wb.save(out_docs)
    shutil.copy(out_docs, os.path.join(ROOT, "docs", "GLOBAL_LOCK_DATA_INDEX.xlsx"))
    shutil.copy(out_docs, os.path.join("/home/user", "03_GLOBAL_LOCK_DATA_INDEX.xlsx"))
    print("Master index Excel refreshed with Electrical, Insurance & Sensor specs!")

if __name__ == "__main__":
    build_master_excel()
